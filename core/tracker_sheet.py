"""The user's "Job application tracker" as the first sheet of the Excel board.

Layout, headers, summary tiles and formulas are taken from the template the user
supplied (Downloads/Job application tracker.xlsx). Rows are filled from every
application that reached `applied`, refreshed on each board rebuild.

The whole point is that this sheet is **edited by hand** — Status moves to
Interview or Rejected, notes and contacts get typed in — while the board around
it is regenerated from scratch every day. So a rebuild must never clobber what
the user wrote. Rows are keyed by `platform:external_id` in a hidden Key column
and merged: columns the user owns are carried forward, columns the pipeline owns
are refreshed, and genuinely new applications are appended.

Formulas use plain A1 ranges rather than the template's `tbl_Tracker[...]`
structured references. They compute exactly the same thing, but structured
references do not survive the LibreOffice recalculation used to verify the
workbook, and a formula that cannot be evaluated ships as a literal `#NAME?`.
"""
from __future__ import annotations

import sqlite3
from datetime import datetime
from pathlib import Path

from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

SHEET_NAME = "Tracker"
TABLE_NAME = "tbl_Tracker"

HEADERS = ["Position", "Company", "Industry", "Location", "Contact",
           "Salary low", "Salary high", "Type", "Schedule", "Applied",
           "Days", "Method", "Status", "Follow-up", "Followed up", "Notes",
           "Key"]
HEADER_ROW = 7
FIRST_DATA_ROW = HEADER_ROW + 1
KEY_COL = len(HEADERS)                      # hidden, holds platform:external_id

# Columns the user owns. Whatever is in the sheet wins over anything we'd
# generate, so hand edits survive every rebuild.
USER_COLUMNS = {"Industry", "Contact", "Type", "Schedule", "Status",
                "Followed up", "Notes"}

WIDTHS = {"A": 34, "B": 26, "C": 16, "D": 22, "E": 24, "F": 12, "G": 12,
          "H": 13, "I": 12, "J": 12, "K": 7.8, "L": 11, "M": 13.1, "N": 12.3,
          "O": 13, "P": 40}

BODY_FONT = Font(name="Aptos Narrow", size=12)
HEAD_FONT = Font(name="Aptos Narrow", size=12, bold=True)
TILE_FONT = Font(name="Aptos Narrow", size=28, bold=True)
LABEL_FONT = Font(name="Aptos Narrow", size=11, color="595959")
DATE_FMT = "yyyy\\-mm\\-dd;@"
LPA_FMT = '0.0" LPA";;"-"'                  # rupees per annum, not dollars

VALIDATIONS = {
    "H": '"Permanent,Fixed term,Temporary,Contract"',
    "I": '"Full-time,Part-time,Freelance"',
    "M": '"Applied,Interview,Rejected,Offer,Withdrawn"',
}


def _applied_rows(db_path: Path) -> list[dict]:
    """Every application that was actually submitted, oldest first.

    Prefers `submissions_log.ts` for the applied date — that is written at the
    moment of submission, whereas `applications.updated_at` moves again if the
    row is touched later.
    """
    con = sqlite3.connect(f"file:{db_path}?mode=ro", uri=True)
    con.row_factory = sqlite3.Row
    try:
        rows = con.execute(
            """SELECT j.platform, j.external_id, j.title, j.company,
                      j.location, j.url, j.salary,
                      COALESCE(MIN(s.ts), a.updated_at) AS applied_at
                 FROM applications a
                 JOIN jobs j ON j.id = a.job_id
                 LEFT JOIN submissions_log s ON s.app_id = a.id
                WHERE a.state = 'applied'
             GROUP BY a.id
             ORDER BY applied_at""").fetchall()
    finally:
        con.close()
    return [dict(r) for r in rows]


def _salary_range(text: str) -> tuple[float | None, float | None]:
    from core.matching import parse_salary_lpa
    parsed = parse_salary_lpa(text or "")
    return parsed if parsed else (None, None)


def _read_existing(ws) -> dict[str, dict]:
    """Key -> the user-owned values currently in the sheet."""
    if ws is None:
        return {}
    index = {h: i for i, h in enumerate(HEADERS)}
    kept: dict[str, dict] = {}
    for row in ws.iter_rows(min_row=FIRST_DATA_ROW, values_only=True):
        if not row or len(row) < KEY_COL:
            continue
        key = row[KEY_COL - 1]
        if not key:
            continue
        kept[str(key)] = {h: row[index[h]] for h in USER_COLUMNS
                          if index[h] < len(row)}
    return kept


def _sync_outcomes(existing: dict[str, dict], db_path: Path) -> int:
    """Push hand-typed Status values back into SQLite.

    The sheet is where the user records what happened after applying, and until
    that reaches the database nothing else can report on it — the pipeline
    stopped at `applied`, so 101 applications carried no outcome at all. Status
    "Applied" is the default we wrote, not a user judgement, so only the other
    values are recorded.
    """
    from core.queue import Queue

    queue = Queue(db_path)
    changed = 0
    for key, vals in existing.items():
        platform, _, external_id = key.partition(":")
        # The "Followed up" checkbox: carried into SQLite so `followups` and
        # the daily report stop nagging about it. Not counted as an outcome.
        if vals.get("Followed up") is True and platform and external_id:
            queue.set_followed_up(platform, external_id)
        status = str(vals.get("Status") or "").strip()
        if not status or status.lower() == "applied":
            continue
        if platform and external_id and queue.set_outcome(
                platform, external_id, status.lower()):
            changed += 1
    return changed


def refresh(workbook, db_path: Path) -> dict:
    """Rebuild the Tracker sheet in `workbook`, preserving hand-entered data."""
    existing = _read_existing(workbook[SHEET_NAME]
                              if SHEET_NAME in workbook.sheetnames else None)
    outcomes_recorded = _sync_outcomes(existing, db_path)
    if SHEET_NAME in workbook.sheetnames:
        del workbook[SHEET_NAME]
    ws = workbook.create_sheet(SHEET_NAME, 0)          # first sheet
    ws.sheet_view.showGridLines = False

    jobs = _applied_rows(db_path)

    # Summary tiles, mirroring the template's four counters.
    last = max(FIRST_DATA_ROW, HEADER_ROW + len(jobs))
    ws["A1"] = f"=COUNTA(A{FIRST_DATA_ROW}:A{last})"
    ws["B1"] = (f'=COUNTIF(M{FIRST_DATA_ROW}:M{last},"Applied")'
                f'+COUNTIF(M{FIRST_DATA_ROW}:M{last},"Interview")'
                f'+COUNTIF(M{FIRST_DATA_ROW}:M{last},"Offer")')
    ws["A4"] = f'=COUNTIF(M{FIRST_DATA_ROW}:M{last},"Interview")'
    ws["B4"] = (f'=COUNTIFS(N{FIRST_DATA_ROW}:N{last},"<="&TODAY(),'
                f'N{FIRST_DATA_ROW}:N{last},">0")')
    ws["A2"], ws["B2"] = "Applications", "Active"
    ws["A5"], ws["B5"] = "Interviews", "Urgent follow-ups"
    for cell in ("A1", "B1", "A4", "B4"):
        ws[cell].font = TILE_FONT
    for cell in ("A2", "B2", "A5", "B5"):
        ws[cell].font = LABEL_FONT

    for col, head in enumerate(HEADERS, start=1):
        c = ws.cell(row=HEADER_ROW, column=col, value=head)
        c.font = HEAD_FONT
        c.alignment = Alignment(vertical="center")
        c.fill = PatternFill("solid", fgColor="F2F2F2")

    carried = 0
    for offset, job in enumerate(jobs):
        r = FIRST_DATA_ROW + offset
        key = f"{job['platform']}:{job['external_id']}"
        prev = existing.get(key, {})
        if prev:
            carried += 1
        lo, hi = _salary_range(job["salary"])
        applied = (job["applied_at"] or "")[:10]

        ws.cell(row=r, column=1, value=job["title"])
        if job["url"]:
            ws.cell(row=r, column=1).hyperlink = job["url"]
            ws.cell(row=r, column=1).style = "Hyperlink"
        ws.cell(row=r, column=2, value=job["company"])
        ws.cell(row=r, column=3, value=prev.get("Industry"))
        ws.cell(row=r, column=4, value=job["location"])
        ws.cell(row=r, column=5, value=prev.get("Contact"))
        ws.cell(row=r, column=6, value=lo).number_format = LPA_FMT
        ws.cell(row=r, column=7, value=hi).number_format = LPA_FMT
        ws.cell(row=r, column=8, value=prev.get("Type"))
        ws.cell(row=r, column=9, value=prev.get("Schedule"))
        applied_cell = ws.cell(row=r, column=10)
        if applied:
            applied_cell.value = datetime.strptime(applied, "%Y-%m-%d")
        applied_cell.number_format = DATE_FMT
        ws.cell(row=r, column=11,
                value=f'=IF(OR(M{r}="Applied",M{r}="Interview"),'
                      f'IF(J{r}="","",TODAY()-J{r}),"")')
        ws.cell(row=r, column=12, value=job["platform"])
        ws.cell(row=r, column=13, value=prev.get("Status") or "Applied")
        ws.cell(row=r, column=14,
                value=f'=IF(M{r}="Applied",IF(O{r},"Done",'
                      f'IF(J{r}="","",J{r}+7)),"-")').number_format = DATE_FMT
        ws.cell(row=r, column=15,
                value=prev.get("Followed up") if prev.get("Followed up") is not None
                else False)
        ws.cell(row=r, column=16, value=prev.get("Notes"))
        ws.cell(row=r, column=KEY_COL, value=key)
        for col in range(1, KEY_COL + 1):
            if ws.cell(row=r, column=col).font.name != "Aptos Narrow":
                ws.cell(row=r, column=col).font = BODY_FONT

    for letter, width in WIDTHS.items():
        ws.column_dimensions[letter].width = width
    ws.column_dimensions[get_column_letter(KEY_COL)].hidden = True
    ws.freeze_panes = f"C{FIRST_DATA_ROW}"

    if jobs:
        ref = f"A{HEADER_ROW}:{get_column_letter(KEY_COL)}{HEADER_ROW + len(jobs)}"
        table = Table(displayName=TABLE_NAME, ref=ref)
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(table)
        for col, formula in VALIDATIONS.items():
            dv = DataValidation(type="list", formula1=formula, allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(f"{col}{FIRST_DATA_ROW}:{col}{HEADER_ROW + len(jobs)}")

    return {"rows": len(jobs), "carried_over": carried,
            "outcomes_recorded": outcomes_recorded}
